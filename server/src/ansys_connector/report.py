"""Engineering report renderer: structured study content -> calc-sheet HTML and Word documents.

Claude supplies the content (it did the engineering); this module only lays it out
consistently, embeds the images and writes the files.
"""

from __future__ import annotations

import base64
import datetime as dt
import html
import mimetypes
import re
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field

from .common import default_workdir, tool

Status = Literal["PASS", "WARN", "FAIL"]


class Row(BaseModel):
    """A labelled value, e.g. a material property, mesh statistic or result."""

    item: str
    value: str
    unit: str = ""
    note: str = ""


class Result(BaseModel):
    quantity: str
    value: str
    unit: str = ""
    location: str = ""
    limit: str = Field("", description="Allowable or target, e.g. '276 MPa yield' or '>= 150 Hz'")
    margin: str = Field("", description="e.g. 'SF 2.0' or '+29 %'")
    status: Status | Literal[""] = ""


class Check(BaseModel):
    check: str = Field(description="What was verified, e.g. 'Equilibrium: reactions vs applied load'")
    expected: str = ""
    actual: str = ""
    status: Status
    note: str = ""


class ConvergenceRow(BaseModel):
    mesh: str
    nodes: str = ""
    value: str
    change: str = Field("", description="Change vs previous mesh, e.g. '0.5 %'")


class ImageRef(BaseModel):
    path: str
    caption: str = ""


class Section(BaseModel):
    heading: str
    body: str = Field("", description="Paragraphs separated by blank lines; lines starting with '- ' become bullets")
    table: list[Row] = []


class ReportSpec(BaseModel):
    title: str = Field(description="What was analysed, e.g. 'Motor mounting bracket - static and modal check'")
    project: str = ""
    part: str = ""
    analyst: str = "Claude via Ansys connector"
    revision: str = "A"
    units: str = "mm, N, MPa, s"
    verdict: Literal["PASS", "FAIL", "REVIEW"] = Field(
        description="PASS only if every requirement is met AND verification passed; REVIEW when results are "
        "uncertain or an assumption needs confirming"
    )
    summary: str = Field(description="3-6 sentence executive summary: what was checked, the key numbers, the verdict")
    scope: str = ""
    geometry: str = ""
    material: list[Row] = []
    mesh: list[Row] = []
    supports: list[str] = []
    loads: list[str] = []
    results: list[Result] = []
    verification: list[Check] = []
    convergence: list[ConvergenceRow] = []
    convergence_quantity: str = ""
    images: list[ImageRef] = []
    assumptions: list[str] = []
    limitations: list[str] = []
    conclusions: list[str] = []
    recommendations: list[str] = []
    extra_sections: list[Section] = []
    files: list[Row] = Field([], description="Model/result files: item = label, value = path")
    sign_off: bool = Field(True, description="Add 'Checked by / Approved by' lines for a human engineer")


# --------------------------------------------------------------------------- HTML

_CSS = """
:root { --ink:#1d252c; --soft:#5b6670; --rule:#cfd6dc; --head:#eef2f5; --accent:#0f5f6b;
  --pass:#1f7a3d; --pass-bg:#e5f4ea; --warn:#8a5a00; --warn-bg:#fdf1db; --fail:#a12626; --fail-bg:#fbe5e5;
  color-scheme: light; }
* { box-sizing:border-box; }
body { margin:0; background:#e9edf0; color:var(--ink); font:14px/1.5 "Segoe UI", system-ui, -apple-system, Roboto, sans-serif; }
.sheet { max-width:860px; margin:24px auto; background:#fff; padding:36px 44px 44px; box-shadow:0 1px 4px rgba(0,0,0,.12); }
h1 { font-size:24px; margin:0 0 4px; text-wrap:balance; }
h2 { font-size:16px; margin:28px 0 10px; padding-bottom:4px; border-bottom:2px solid var(--ink); text-transform:uppercase; letter-spacing:.04em; }
h3 { font-size:14px; margin:16px 0 6px; }
p { margin:0 0 8px; max-width:72ch; }
table { width:100%; border-collapse:collapse; margin:6px 0 10px; font-variant-numeric:tabular-nums; }
th, td { border:1px solid var(--rule); padding:5px 8px; text-align:left; vertical-align:top; }
th { background:var(--head); font-weight:600; font-size:12px; text-transform:uppercase; letter-spacing:.03em; }
.tb { display:grid; grid-template-columns:1fr 230px; border:2px solid var(--ink); margin-bottom:18px; }
.tb-main { padding:14px 16px; border-right:2px solid var(--ink); }
.tb-meta { display:grid; font-size:12px; }
.tb-meta div { padding:4px 10px; border-bottom:1px solid var(--rule); }
.tb-meta div:last-child { border-bottom:0; }
.tb-meta span { display:block; color:var(--soft); font-size:10px; text-transform:uppercase; letter-spacing:.08em; }
.verdict { display:inline-block; font-weight:700; padding:3px 12px; border-radius:3px; letter-spacing:.06em; }
.chip { display:inline-block; font-weight:700; font-size:11px; padding:1px 7px; border-radius:3px; }
.PASS { color:var(--pass); background:var(--pass-bg); } .WARN, .REVIEW { color:var(--warn); background:var(--warn-bg); }
.FAIL { color:var(--fail); background:var(--fail-bg); }
.summary { border-left:4px solid var(--accent); background:#f4f8f9; padding:10px 14px; margin:8px 0 4px; }
figure { margin:12px 0; } figure img { max-width:100%; border:1px solid var(--rule); }
figcaption { color:var(--soft); font-size:12px; margin-top:4px; }
ul { margin:4px 0 10px; padding-left:20px; } li { margin:2px 0; }
.signoff { display:grid; grid-template-columns:repeat(3,1fr); gap:18px; margin-top:12px; }
.signoff div { border-top:1px solid var(--ink); padding-top:4px; font-size:12px; color:var(--soft); }
.foot { margin-top:28px; color:var(--soft); font-size:11px; border-top:1px solid var(--rule); padding-top:8px; }
@media (max-width:640px) { .sheet { padding:20px 16px; margin:0; } .tb { grid-template-columns:1fr; } .tb-main { border-right:0; border-bottom:2px solid var(--ink); } .signoff { grid-template-columns:1fr; } }
@media print { body { background:#fff; } .sheet { box-shadow:none; margin:0; max-width:none; padding:0; } h2 { break-after:avoid; } figure, table { break-inside:avoid; } }
"""


def _e(text: str) -> str:
    return html.escape(str(text))


def _paragraphs(body: str) -> str:
    out = []
    for block in re.split(r"\n\s*\n", body.strip()):
        lines = block.splitlines()
        if lines and all(ln.lstrip().startswith("- ") for ln in lines):
            out.append("<ul>" + "".join(f"<li>{_e(ln.lstrip()[2:])}</li>" for ln in lines) + "</ul>")
        elif block.strip():
            out.append(f"<p>{_e(' '.join(ln.strip() for ln in lines))}</p>")
    return "".join(out)


def _table(headers: list[str], rows: list[list[str]], status_col: int | None = None) -> str:
    head = "".join(f"<th>{_e(h)}</th>" for h in headers)
    body = []
    for row in rows:
        cells = []
        for i, cell in enumerate(row):
            if i == status_col and cell:
                cells.append(f'<td><span class="chip {_e(cell)}">{_e(cell)}</span></td>')
            else:
                cells.append(f"<td>{_e(cell)}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _bullets(items: list[str]) -> str:
    return "<ul>" + "".join(f"<li>{_e(i)}</li>" for i in items) + "</ul>"


def _img_data_uri(path: Path) -> str:
    mime = mimetypes.guess_type(path.name)[0] or "image/png"
    return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode()}"


def _rows(rows: list[Row]) -> list[list[str]]:
    return [[r.item, r.value, r.unit, r.note] for r in rows]


def render_html(spec: ReportSpec, today: str) -> str:
    s = spec
    parts = [
        f"<title>{_e(s.title)}</title><style>{_CSS}</style><div class='sheet'>",
        "<div class='tb'><div class='tb-main'>",
        f"<h1>{_e(s.title)}</h1>",
        f"<p>{_e(' / '.join(x for x in (s.project, s.part) if x))}</p>" if (s.project or s.part) else "",
        f"<span class='verdict {s.verdict}'>{s.verdict}</span></div><div class='tb-meta'>",
        f"<div><span>Analyst</span>{_e(s.analyst)}</div><div><span>Date</span>{today}</div>",
        f"<div><span>Revision</span>{_e(s.revision)}</div><div><span>Units</span>{_e(s.units)}</div></div></div>",
        f"<div class='summary'>{_paragraphs(s.summary)}</div>",
    ]
    if s.scope:
        parts += ["<h2>1. Scope</h2>", _paragraphs(s.scope)]
    parts.append("<h2>2. Model</h2>")
    if s.geometry:
        parts += ["<h3>Geometry</h3>", _paragraphs(s.geometry)]
    if s.material:
        parts += ["<h3>Material</h3>", _table(["Property", "Value", "Unit", "Note"], _rows(s.material))]
    if s.mesh:
        parts += ["<h3>Mesh</h3>", _table(["Item", "Value", "Unit", "Note"], _rows(s.mesh))]
    if s.supports:
        parts += ["<h3>Supports</h3>", _bullets(s.supports)]
    if s.loads:
        parts += ["<h3>Loads</h3>", _bullets(s.loads)]
    if s.results:
        parts += [
            "<h2>3. Results</h2>",
            _table(
                ["Quantity", "Value", "Unit", "Location", "Limit", "Margin", "Status"],
                [[r.quantity, r.value, r.unit, r.location, r.limit, r.margin, r.status] for r in s.results],
                status_col=6,
            ),
        ]
    for img in s.images:
        p = Path(img.path)
        if p.is_file():
            parts.append(f"<figure><img src='{_img_data_uri(p)}' alt='{_e(img.caption)}'><figcaption>{_e(img.caption)}</figcaption></figure>")
    if s.verification or s.convergence:
        parts.append("<h2>4. Verification</h2>")
    if s.verification:
        parts.append(
            _table(["Check", "Expected", "Actual", "Status", "Note"],
                   [[c.check, c.expected, c.actual, c.status, c.note] for c in s.verification], status_col=3)
        )
    if s.convergence:
        parts += [
            f"<h3>Mesh convergence{': ' + _e(s.convergence_quantity) if s.convergence_quantity else ''}</h3>",
            _table(["Mesh", "Nodes", "Value", "Change"], [[c.mesh, c.nodes, c.value, c.change] for c in s.convergence]),
        ]
    n = 5
    for heading, items in (("Assumptions", s.assumptions), ("Limitations", s.limitations),
                           ("Conclusions", s.conclusions), ("Recommendations", s.recommendations)):
        if items:
            parts += [f"<h2>{n}. {heading}</h2>", _bullets(items)]
            n += 1
    for sec in s.extra_sections:
        parts += [f"<h2>{n}. {_e(sec.heading)}</h2>", _paragraphs(sec.body)]
        if sec.table:
            parts.append(_table(["Item", "Value", "Unit", "Note"], _rows(sec.table)))
        n += 1
    if s.files:
        parts += [f"<h2>{n}. Files</h2>", _table(["File", "Path", "", ""], _rows(s.files))]
    if s.sign_off:
        parts.append("<div class='signoff'><div>Prepared: " + _e(s.analyst) + "</div><div>Checked by / date</div><div>Approved by / date</div></div>")
    parts.append(
        "<p class='foot'>Generated with the Ansys connector for Claude (unofficial). Simulation results depend on the "
        "model, inputs and assumptions listed above and must be reviewed by a qualified engineer before use.</p></div>"
    )
    return "".join(parts)


# --------------------------------------------------------------------------- Word

def render_docx(spec: ReportSpec, today: str, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    s = spec
    colors = {"PASS": RGBColor(0x1F, 0x7A, 0x3D), "WARN": RGBColor(0x8A, 0x5A, 0x00),
              "REVIEW": RGBColor(0x8A, 0x5A, 0x00), "FAIL": RGBColor(0xA1, 0x26, 0x26)}
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    def table(headers, rows, status_col=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            run = t.rows[0].cells[i].paragraphs[0].add_run(h)
            run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                run = cells[i].paragraphs[0].add_run(str(val))
                if i == status_col and val in colors:
                    run.bold = True
                    run.font.color.rgb = colors[val]
        doc.add_paragraph()

    doc.add_heading(s.title, 0)
    meta = doc.add_paragraph()
    meta.add_run(" / ".join(x for x in (s.project, s.part) if x) + "\n" if (s.project or s.part) else "")
    meta.add_run(f"Analyst: {s.analyst}    Date: {today}    Rev: {s.revision}    Units: {s.units}\n")
    v = meta.add_run(f"Verdict: {s.verdict}")
    v.bold = True
    v.font.color.rgb = colors.get(s.verdict, RGBColor(0, 0, 0))
    doc.add_paragraph(s.summary)

    if s.scope:
        doc.add_heading("1. Scope", 1)
        doc.add_paragraph(s.scope)
    doc.add_heading("2. Model", 1)
    if s.geometry:
        doc.add_paragraph(s.geometry)
    if s.material:
        doc.add_heading("Material", 2)
        table(["Property", "Value", "Unit", "Note"], _rows(s.material))
    if s.mesh:
        doc.add_heading("Mesh", 2)
        table(["Item", "Value", "Unit", "Note"], _rows(s.mesh))
    for heading, items in (("Supports", s.supports), ("Loads", s.loads)):
        if items:
            doc.add_heading(heading, 2)
            for it in items:
                doc.add_paragraph(it, style="List Bullet")
    if s.results:
        doc.add_heading("3. Results", 1)
        table(["Quantity", "Value", "Unit", "Location", "Limit", "Margin", "Status"],
              [[r.quantity, r.value, r.unit, r.location, r.limit, r.margin, r.status] for r in s.results], status_col=6)
    for img in s.images:
        p = Path(img.path)
        if p.is_file():
            doc.add_picture(str(p), width=Inches(6.0))
            cap = doc.add_paragraph(img.caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if s.verification or s.convergence:
        doc.add_heading("4. Verification", 1)
    if s.verification:
        table(["Check", "Expected", "Actual", "Status", "Note"],
              [[c.check, c.expected, c.actual, c.status, c.note] for c in s.verification], status_col=3)
    if s.convergence:
        doc.add_heading("Mesh convergence" + (f": {s.convergence_quantity}" if s.convergence_quantity else ""), 2)
        table(["Mesh", "Nodes", "Value", "Change"], [[c.mesh, c.nodes, c.value, c.change] for c in s.convergence])
    n = 5
    for heading, items in (("Assumptions", s.assumptions), ("Limitations", s.limitations),
                           ("Conclusions", s.conclusions), ("Recommendations", s.recommendations)):
        if items:
            doc.add_heading(f"{n}. {heading}", 1)
            for it in items:
                doc.add_paragraph(it, style="List Bullet")
            n += 1
    for sec in s.extra_sections:
        doc.add_heading(f"{n}. {sec.heading}", 1)
        for block in re.split(r"\n\s*\n", sec.body.strip()):
            for ln in block.splitlines():
                if ln.lstrip().startswith("- "):
                    doc.add_paragraph(ln.lstrip()[2:], style="List Bullet")
            if not all(ln.lstrip().startswith("- ") for ln in block.splitlines()) and block.strip():
                doc.add_paragraph(" ".join(ln.strip() for ln in block.splitlines()))
        if sec.table:
            table(["Item", "Value", "Unit", "Note"], _rows(sec.table))
        n += 1
    if s.files:
        doc.add_heading(f"{n}. Files", 1)
        table(["File", "Path", "", ""], _rows(s.files))
    if s.sign_off:
        doc.add_paragraph("\nPrepared: " + s.analyst + "\n\nChecked by: ____________________   Date: ________\n\n"
                          "Approved by: ___________________   Date: ________")
    foot = doc.add_paragraph("Generated with the Ansys connector for Claude (unofficial). Results depend on the model, "
                             "inputs and assumptions above and must be reviewed by a qualified engineer before use.")
    foot.runs[0].font.size = Pt(8)
    doc.save(str(path))


@tool
def report_create(spec: ReportSpec, output_dir: str | None = None, filename: str | None = None,
                  formats: list[Literal["html", "docx"]] = ["html", "docx"]) -> dict:
    """Write an engineering report for a finished study as HTML (self-contained, printable) and/or Word.

    Fill `spec` from results you have ALREADY verified: key results with limits and margins, the
    verification checks you actually ran (status PASS/WARN/FAIL - never mark a check you did not
    run), the mesh-convergence table, image paths from the solver tools, assumptions and limitations.
    Returns the file paths; nothing is re-computed here.
    """
    folder = Path(output_dir) if output_dir else default_workdir("reports")
    folder.mkdir(parents=True, exist_ok=True)
    stem = filename or re.sub(r"[^A-Za-z0-9_-]+", "_", spec.title).strip("_")[:60] or "report"
    today = dt.date.today().isoformat()
    written = {}
    missing = [i.path for i in spec.images if not Path(i.path).is_file()]
    if "html" in formats:
        p = folder / f"{stem}.html"
        p.write_text(render_html(spec, today), encoding="utf-8")
        written["html"] = str(p)
    if "docx" in formats:
        p = folder / f"{stem}.docx"
        render_docx(spec, today, p)
        written["docx"] = str(p)
    out = {"files": written, "verdict": spec.verdict,
           "checks": {k: sum(c.status == k for c in spec.verification) for k in ("PASS", "WARN", "FAIL")}}
    if missing:
        out["missing_images"] = missing
    return out
